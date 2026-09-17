import streamlit as st, pandas as pd, os, io, zipfile, requests, oracledb

W_DIR = os.path.join(os.getcwd(), "wallet_files")
SECRETS_FILE_PATH = os.path.join(os.getcwd(), ".streamlit", "secrets.toml")

def get_db_connection():
    is_streamlit_active = False
    try:
        if st.runtime.exists(): is_streamlit_active = True
    except: pass
    
    if is_streamlit_active:
        # 🌐 Web Portal Mode: Inherits keys cleanly from Streamlit's runtime memory secrets
        db_pwd = st.secrets["db_password"]
        wallet_pwd = st.secrets["wallet_password"]
    else:
        # 💻 Bare Console Script Mode: Parses your physical secrets.toml file directly from your disk folder!
        db_pwd, wallet_pwd = None, None
        if os.path.exists(SECRETS_FILE_PATH):
            try:
                # Custom lightweight parser block to read TOML key-value rows safely without extra libraries
                with open(SECRETS_FILE_PATH, "r", encoding="utf-8") as f_sec:
                    for line in f_sec:
                        clean_line = line.strip()
                        if "=" in clean_line and not clean_line.startswith("#"):
                            key, val = clean_line.split("=", 1)
                            key_cl = key.strip().strip('"').strip("'")
                            val_cl = val.strip().strip('"').strip("'")
                            if key_cl == "db_password": db_pwd = val_cl
                            elif key_cl == "wallet_password": wallet_pwd = val_cl
            except Exception as e:
                print(f"⚠️ Warning: Could not read local secrets.toml file: {e}")
        
        # Fallback security check guard
        if not db_pwd or not wallet_pwd:
            raise ValueError(f"❌ Error: Secure credentials missing! Ensure keys are active inside '{SECRETS_FILE_PATH}'")

    p = {
        "user": "ADMIN", "password": db_pwd, "dsn": "search_low",
        "config_dir": W_DIR, "wallet_location": W_DIR, "wallet_password": wallet_pwd,
        "ssl_server_dn_match": False
    }
    def bh(cursor, name, dtype, size, prec, scale):
        if dtype == oracledb.DB_TYPE_BLOB: return cursor.var(bytes, arraysize=cursor.arraysize)
    conn = oracledb.connect(**p)
    conn.outputtypehandler = bh
    return conn

def get_master_taxonomy():
    return {
        "Delivery Manager": {"☁️ Cloud": ["OCI", "AWS", "Azure"], "📐 Design": ["Database Modelling", "Oracle Designer"], "🤖 AI": ["AI Skills"]},
        "Database Administrator (DBA)": {"💎 Core": ["Oracle DBA", "PL/SQL", "SQL"], "💾 Open": ["PostgreSQL"]},
        "Application Developer": {"🐍 Python": ["Python", "Flask", "Django", "Streamlit"], "☕ Java": ["Java"], "🤖 AI": ["AI Skills"]},
        "System Administrator": {"💿 Unix": ["Linux", "Unix"], "⚙️ Legacy": ["Websphere", "Dos", "Mac"]},
        "ITIL & Infrastructure Engineer": {"📦 DevOps": ["Docker", "Kubernetes", "Git"], "🧱 Legacy": ["Cobol"]}
    }

def query_matched_profiles_via_stored_function(keyword=None, location=None):
    is_cloud = os.environ.get("STREAMLIT_RUNTIME_ENV") or "mount" in os.getcwd()
    
    if isinstance(keyword, (list, tuple, set)): kw_param = ",".join([str(k) for k in keyword]).strip()
    else: kw_param = str(keyword).strip() if keyword else ""
        
    if isinstance(location, (list, tuple, set)): loc_param = ",".join([str(l) for l in location]).strip()
    else: loc_param = str(location).strip() if location else ""
        
    if kw_param.lower() == "none" or kw_param == "": kw_param = ""
    if loc_param.lower() == "none" or loc_param == "": loc_param = ""
    
    if is_cloud:
        try:
            conn = get_db_connection()
            cursor = conn.cursor()
            ref_cursor = cursor.callfunc(
                "GET_MATCHED_CANDIDATES", oracledb.DB_TYPE_CURSOR, 
                keyword_parameters={"P_KEYWORD": kw_param if kw_param else None, "P_LOCATION": loc_param if loc_param else None}
            )
            rows = ref_cursor.fetchall()
            cols = [col.name.upper() for col in ref_cursor.description]
            results = []
            for row in rows:
                record = dict(zip(cols, row))
                record["HAS_BLOB"] = True if "RESUME_BLOB" in record and record["RESUME_BLOB"] is not None else False
                results.append(record)
            ref_cursor.close(); cursor.close(); conn.close()
            df = pd.DataFrame(results)
            if not df.empty: df.columns = [c.upper() for c in df.columns]
            return df
        except Exception as cloud_err:
            st.error(f"Cloud Direct DB Error: {cloud_err}"); return pd.DataFrame()
    else:
        try:
            api_url = "http://localhost:8000/api/candidates"
            payload_params = {}
            if kw_param: payload_params["keyword"] = kw_param
            if loc_param: payload_params["location"] = loc_param
            response = requests.get(api_url, params=payload_params, timeout=5)
            if response.status_code == 200:
                json_data = response.json()
                if json_data.get("status") == "SUCCESS":
                    df = pd.DataFrame(json_data.get("data", []))
                    if not df.empty: df.columns = [c.upper() for c in df.columns]
                    return df
            return pd.DataFrame()
        except Exception as local_err:
            st.error(f"⚠️ Local Microservice Connection Failure: Ensure api_server.py is running via Uvicorn!"); return pd.DataFrame()

def build_zip_archive(candidates):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for c in candidates:
            if c.get("BLOB") is not None: zf.writestr(f"{c['NAME'].replace(' ', '_')}_Resume.pdf", c["BLOB"])
    buf.seek(0); return buf.getvalue()

def init_page_headers():
    st.set_page_config(page_title="Talent Search", layout="wide")
    st.title("☁️ Talent Search Workspace")
    st.write("Tabular workstation powered by Oracle Cloud Infrastructure.")
    st.markdown("<hr>", unsafe_allow_html=True)

def apply_corporate_styles(wb):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    f_hdr = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    f_cel = Font(name="Segoe UI", size=10, color="2C3E50")
    fill_hdr = PatternFill(start_color="1F77B4", end_color="1F77B4", fill_type="solid")
    fill_zb = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")
    fill_wh = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    border_thin = Border(left=Side(style="thin", color="E0E0E0"), right=Side(style="thin", color="E0E0E0"), top=Side(style="thin", color="E0E0E0"), bottom=Side(style="thin", color="E0E0E0"))
    for name in wb.sheetnames:
        ws = wb[name]; ws.sheet_view.showGridLines = True
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = 26 if r == 1 else 22
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                if r == 1: cell.font = f_hdr; cell.fill = fill_hdr
                else:
                    cell.font = f_cel; cell.border = border_thin
                    cell.fill = fill_zb if r % 2 == 0 else fill_wh
                    if str(cell.value).startswith("Go to") or str(cell.value).startswith("2026-"): cell.alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in range(1, ws.max_column + 1):
            lens = [len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(1, ws.max_row + 1)]
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max((max(lens) if lens else 10) + 3, 12), 75)
